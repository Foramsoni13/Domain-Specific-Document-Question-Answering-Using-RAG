"""
Document Loader and Preprocessing Module for Domain-Specific RAG.
Supports PDF, TXT, MD, and DOCX files with page number tracking,
metadata preservation, and NLP text normalization.
"""

import os
import re
import unicodedata
from typing import List, Dict, Any, Optional

class DocumentPage:
    """Represents a single page or logical unit of a document with metadata."""
    def __init__(self, doc_name: str, page_number: int, content: str, metadata: Optional[Dict[str, Any]] = None):
        self.doc_name = doc_name
        self.page_number = page_number
        self.content = content
        self.metadata = metadata or {}

    def to_dict(self) -> Dict[str, Any]:
        return {
            "doc_name": self.doc_name,
            "page_number": self.page_number,
            "content": self.content,
            "metadata": self.metadata
        }

    def __repr__(self) -> str:
        return f"<DocumentPage doc='{self.doc_name}' page={self.page_number} chars={len(self.content)}>"


class DocumentPreprocessor:
    """Performs domain-aware text cleaning and normalization."""

    @staticmethod
    def clean_text(text: str) -> str:
        """
        Cleans and normalizes text while preserving domain-critical structures
        such as bullet points, section numbering (e.g. 1.2), percentages, and currency.
        """
        if not text:
            return ""
        
        # 1. Unicode normalization (NFKC)
        text = unicodedata.normalize("NFKC", text)

        # 2. Standardize quotes and hyphens
        text = text.replace("“", "\"").replace("”", "\"").replace("‘", "'").replace("’", "'")
        text = text.replace("—", " - ").replace("–", " - ")

        # 3. Replace multiple horizontal whitespace with a single space
        text = re.sub(r"[ \t]+", " ", text)

        # 4. Standardize consecutive newlines (max 2 consecutive newlines)
        text = re.sub(r"\n\s*\n", "\n\n", text)

        # 5. Remove trailing/leading spaces on lines
        lines = [line.strip() for line in text.split("\n")]
        text = "\n".join(lines).strip()

        return text


class DocumentLoader:
    """Loads documents across various formats with page-level granularity."""

    def __init__(self, preprocessor: Optional[DocumentPreprocessor] = None):
        self.preprocessor = preprocessor or DocumentPreprocessor()

    def load_pdf(self, file_path: str) -> List[DocumentPage]:
        """Loads a PDF file and extracts text page-by-page."""
        pages: List[DocumentPage] = []
        doc_name = os.path.basename(file_path)

        # Attempt extraction via pypdfium2 first (fast, accurate layout parsing)
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(file_path)
            try:
                for page_idx, page in enumerate(pdf):
                    text_page = page.get_textpage()
                    raw_text = text_page.get_text_range()
                    text_page.close()
                    page.close()
                    clean_text = self.preprocessor.clean_text(raw_text)
                    if clean_text:
                        pages.append(DocumentPage(
                            doc_name=doc_name,
                            page_number=page_idx + 1,
                            content=clean_text,
                            metadata={"file_path": file_path, "total_pages": len(pdf), "source_type": "pdf"}
                        ))
            finally:
                pdf.close()
            if pages:
                return pages
        except Exception as e:
            # Fallback to PyPDF2 if pypdfium2 fails or is unavailable
            pass

        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                num_pages = len(reader.pages)
                for page_idx, page in enumerate(reader.pages):
                    raw_text = page.extract_text() or ""
                    clean_text = self.preprocessor.clean_text(raw_text)
                    if clean_text:
                        pages.append(DocumentPage(
                            doc_name=doc_name,
                            page_number=page_idx + 1,
                            content=clean_text,
                            metadata={"file_path": file_path, "total_pages": num_pages, "source_type": "pdf"}
                        ))
        except Exception as e:
            print(f"[DocumentLoader] Error reading PDF {file_path}: {e}")

        return pages

    def load_text(self, file_path: str) -> List[DocumentPage]:
        """
        Loads plain text or markdown files.
        Detects explicit page markers like `--- [PAGE X] ---` or falls back to single-page.
        """
        doc_name = os.path.basename(file_path)
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            raw_content = f.read()

        pages: List[DocumentPage] = []
        # Check for page markers
        page_splits = re.split(r"---\s*\[PAGE\s+(\d+)\]\s*---", raw_content, flags=re.IGNORECASE)

        if len(page_splits) > 1:
            # First element is pre-first-page header
            header_text = page_splits[0].strip()
            for i in range(1, len(page_splits), 2):
                p_num = int(page_splits[i])
                p_body = page_splits[i + 1].strip()
                full_body = (header_text + "\n\n" + p_body) if (i == 1 and header_text) else p_body
                clean_body = self.preprocessor.clean_text(full_body)
                if clean_body:
                    pages.append(DocumentPage(
                        doc_name=doc_name,
                        page_number=p_num,
                        content=clean_body,
                        metadata={"file_path": file_path, "source_type": "text"}
                    ))
        else:
            clean_content = self.preprocessor.clean_text(raw_content)
            pages.append(DocumentPage(
                doc_name=doc_name,
                page_number=1,
                content=clean_content,
                metadata={"file_path": file_path, "source_type": "text"}
            ))

        return pages

    def load_docx(self, file_path: str) -> List[DocumentPage]:
        """Loads Microsoft Word documents (.docx)."""
        doc_name = os.path.basename(file_path)
        pages: List[DocumentPage] = []
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            clean_text = self.preprocessor.clean_text(full_text)
            pages.append(DocumentPage(
                doc_name=doc_name,
                page_number=1,
                content=clean_text,
                metadata={"file_path": file_path, "source_type": "docx"}
            ))
        except Exception as e:
            print(f"[DocumentLoader] Error loading docx {file_path}: {e}")
        return pages

    def load_file(self, file_path: str) -> List[DocumentPage]:
        """Dispatches file to the appropriate format loader based on extension."""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return self.load_pdf(file_path)
        elif ext in [".txt", ".md", ".log"]:
            return self.load_text(file_path)
        elif ext in [".docx", ".doc"]:
            return self.load_docx(file_path)
        else:
            # Fallback as text
            return self.load_text(file_path)

    def load_directory(self, dir_path: str) -> List[DocumentPage]:
        """Loads all supported documents from a directory."""
        all_pages: List[DocumentPage] = []
        if not os.path.exists(dir_path):
            return all_pages

        for root, _, files in os.walk(dir_path):
            for file in sorted(files):
                if file.startswith("."):
                    continue
                file_path = os.path.join(root, file)
                ext = os.path.splitext(file)[1].lower()
                if ext in [".pdf", ".txt", ".md", ".docx"]:
                    pages = self.load_file(file_path)
                    all_pages.extend(pages)
        return all_pages
